#!/usr/bin/env python3
"""
====================================================================
VISA C2 CHECK AUTOMATION - Python Version
====================================================================
Otomasi pencocokan file PDF Visa C2 di folder lokal
dengan daftar peserta di file Excel/CSV.

CARA PAKAI:
  python3 visa_check_automation.py --folder <path_folder_pdf> --input <file_excel_atau_csv> --output <file_output>

CONTOH:
  python3 visa_check_automation.py --folder ./visa_pdfs --input peserta.xlsx --output hasil.xlsx
  python3 visa_check_automation.py --folder ./visa_pdfs --input peserta.csv --output hasil.csv

STRUKTUR FOLDER PDF:
  visa_pdfs/
  ├── Singapore/
  │   ├── Visa C2 2026 Nama Peserta.pdf
  │   └── ...
  ├── Thailand/
  │   └── ...
  └── Vietnam/
      └── ...

STRUKTUR FILE INPUT (Excel/CSV):
  Kolom: No. | Title | Full Name | Country | Company | Participant Role | Batch | Ket
  (Kolom Ket akan diisi √ jika cocok)
====================================================================
"""

import argparse
import csv
import os
import re
import sys

# ===================== KONFIGURASI =====================

MATCH_THRESHOLD = 0.7
FUZZY_WORD_THRESHOLD = 0.65

# Kolom (0-indexed) - sesuaikan jika format berbeda
COL_NAME = 2       # Kolom C = Full Name
COL_COUNTRY = 3    # Kolom D = Country
COL_KET = 7        # Kolom H = Ket


# ===================== NAME MATCHING =====================

def extract_name_from_pdf(filename):
    name = re.sub(r'\.pdf$', '', filename, flags=re.IGNORECASE)
    name = re.sub(r'^Visa\s+C2\s+\d+\s+', '', name, flags=re.IGNORECASE)
    return name.strip()


def normalize_name(name):
    return re.sub(r'\s+', ' ', name.lower().strip())


def name_words(name):
    return set(normalize_name(name).split())


def levenshtein(s1, s2):
    if len(s1) < len(s2):
        return levenshtein(s2, s1)
    if len(s2) == 0:
        return len(s1)
    prev_row = list(range(len(s2) + 1))
    for i, c1 in enumerate(s1):
        curr_row = [i + 1]
        for j, c2 in enumerate(s2):
            curr_row.append(min(
                prev_row[j + 1] + 1,
                curr_row[j] + 1,
                prev_row[j] + (c1 != c2)
            ))
        prev_row = curr_row
    return prev_row[-1]


def word_similarity(w1, w2):
    if w1 == w2:
        return 1.0
    if w1 in w2 or w2 in w1:
        return max(len(w1), len(w2)) / (max(len(w1), len(w2)) + 1)
    d = levenshtein(w1, w2)
    max_len = max(len(w1), len(w2))
    return 1.0 - d / max_len if max_len > 0 else 0.0


def fuzzy_word_match(words1, words2):
    if not words1 or not words2:
        return 0.0
    common = words1 & words2
    total = words1 | words2
    if len(common) == len(total):
        return 1.0

    unmatched1 = words1 - common
    unmatched2 = words2 - common
    fuzzy_matches = 0
    used2 = set()

    for w1 in unmatched1:
        best_score = 0
        best_w2 = None
        for w2 in unmatched2:
            if w2 in used2:
                continue
            score = word_similarity(w1, w2)
            if score > best_score:
                best_score = score
                best_w2 = w2
        if best_score >= FUZZY_WORD_THRESHOLD and best_w2:
            fuzzy_matches += 1
            used2.add(best_w2)

    exact_weight = len(common) * 1.0
    fuzzy_weight = fuzzy_matches * 0.7
    total_words = max(len(words1), len(words2))
    return (exact_weight + fuzzy_weight) / total_words if total_words > 0 else 0.0


# ===================== FOLDER SCANNING =====================

def scan_pdf_folder(folder_path):
    pdfs_by_country = {}

    if not os.path.isdir(folder_path):
        print(f"ERROR: Folder '{folder_path}' tidak ditemukan!")
        sys.exit(1)

    for entry in os.listdir(folder_path):
        subfolder = os.path.join(folder_path, entry)
        if os.path.isdir(subfolder):
            country = entry.lower()
            names = []
            for f in os.listdir(subfolder):
                if f.lower().endswith('.pdf'):
                    extracted = extract_name_from_pdf(f)
                    if extracted:
                        names.append((extracted, name_words(extracted)))
            if names:
                pdfs_by_country[country] = names

    # Juga cek PDF di root folder
    root_pdfs = []
    for f in os.listdir(folder_path):
        if os.path.isfile(os.path.join(folder_path, f)) and f.lower().endswith('.pdf'):
            extracted = extract_name_from_pdf(f)
            if extracted:
                root_pdfs.append((extracted, name_words(extracted)))
    if root_pdfs:
        pdfs_by_country["_root"] = root_pdfs

    return pdfs_by_country


# ===================== FILE I/O =====================

def read_input(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    rows = []

    if ext == '.csv':
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            header = next(reader)
            for row in reader:
                rows.append(list(row))
        return header, rows

    elif ext in ('.xlsx', '.xls'):
        try:
            import openpyxl
        except ImportError:
            print("ERROR: openpyxl belum terinstall. Jalankan: pip install openpyxl")
            sys.exit(1)
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        header = [str(c or '') for c in all_rows[0]]
        for row in all_rows[1:]:
            rows.append([str(c or '') for c in row])
        return header, rows

    else:
        print(f"ERROR: Format file '{ext}' tidak didukung. Gunakan .csv atau .xlsx")
        sys.exit(1)


def write_output(filepath, header, rows):
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.csv':
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(header)
            writer.writerows(rows)

    elif ext in ('.xlsx', '.xls'):
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, Border, Side
        except ImportError:
            print("ERROR: openpyxl belum terinstall. Jalankan: pip install openpyxl")
            sys.exit(1)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        for col_idx, h in enumerate(header, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        for row_idx, row in enumerate(rows, 2):
            for col_idx, val in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx == COL_KET + 1:
                    cell.alignment = Alignment(horizontal='center')

        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 6
        ws.column_dimensions['C'].width = 45
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 10
        ws.column_dimensions['H'].width = 8

        wb.save(filepath)
    else:
        print(f"ERROR: Format output '{ext}' tidak didukung.")
        sys.exit(1)


# ===================== MAIN =====================

def main():
    parser = argparse.ArgumentParser(
        description="Visa C2 Check Automation - Cocokkan PDF visa dengan daftar peserta"
    )
    parser.add_argument("--folder", required=True, help="Path ke folder yang berisi subfolder negara dengan PDF visa")
    parser.add_argument("--input", required=True, help="File input peserta (.xlsx atau .csv)")
    parser.add_argument("--output", required=True, help="File output hasil (.xlsx atau .csv)")
    parser.add_argument("--threshold", type=float, default=MATCH_THRESHOLD, help=f"Threshold pencocokan (default: {MATCH_THRESHOLD})")
    parser.add_argument("--name-col", type=int, default=COL_NAME, help=f"Index kolom nama (0-based, default: {COL_NAME})")
    parser.add_argument("--country-col", type=int, default=COL_COUNTRY, help=f"Index kolom negara (0-based, default: {COL_COUNTRY})")
    parser.add_argument("--ket-col", type=int, default=COL_KET, help=f"Index kolom Ket (0-based, default: {COL_KET})")

    args = parser.parse_args()

    match_threshold = args.threshold

    print("=" * 60)
    print("VISA C2 CHECK AUTOMATION")
    print("=" * 60)

    # 1. Scan folder PDF
    print(f"\n[1/3] Scanning folder: {args.folder}")
    pdfs_by_country = scan_pdf_folder(args.folder)
    total_pdfs = sum(len(v) for v in pdfs_by_country.values())
    print(f"  Ditemukan {len(pdfs_by_country)} negara, {total_pdfs} file PDF")
    for country, names in sorted(pdfs_by_country.items()):
        if country != "_root":
            print(f"  - {country}: {len(names)} PDF")

    # 2. Baca file input
    print(f"\n[2/3] Membaca file: {args.input}")
    header, rows = read_input(args.input)
    print(f"  {len(rows)} peserta ditemukan")

    # 3. Cocokkan
    print(f"\n[3/3] Mencocokkan nama...")
    matched_count = 0
    unmatched_count = 0
    unmatched_list = []

    for i, row in enumerate(rows):
        # Pastikan row cukup panjang
        while len(row) <= max(args.name_col, args.country_col, args.ket_col):
            row.append('')

        participant_name = row[args.name_col].strip()
        country = row[args.country_col].strip()

        if not participant_name:
            continue

        country_key = country.lower()
        pdf_names = pdfs_by_country.get(country_key, [])

        if not pdf_names:
            unmatched_count += 1
            unmatched_list.append((i + 2, participant_name, country, "Tidak ada folder/PDF"))
            continue

        spreadsheet_words = name_words(participant_name)
        best_score = 0
        best_match = ""

        for pdf_name, pdf_words in pdf_names:
            score = fuzzy_word_match(spreadsheet_words, pdf_words)
            if score > best_score:
                best_score = score
                best_match = pdf_name

        if best_score >= match_threshold:
            matched_count += 1
            row[args.ket_col] = "√"
            print(f"  ✓ Row {i+2}: {participant_name} -> {best_match} ({best_score:.0%})")
        else:
            unmatched_count += 1
            unmatched_list.append((i + 2, participant_name, country, f"Best: {best_score:.0%}"))

    # 4. Simpan output
    write_output(args.output, header, rows)

    # 5. Laporan
    print(f"\n{'=' * 60}")
    print(f"HASIL PENCOCOKAN")
    print(f"{'=' * 60}")
    print(f"  Cocok (√)    : {matched_count} peserta")
    print(f"  Tidak cocok  : {unmatched_count} peserta")
    print(f"  Total        : {matched_count + unmatched_count} peserta")

    if unmatched_list:
        print(f"\nDaftar tidak cocok:")
        for row_num, name, country, reason in unmatched_list:
            print(f"  Row {row_num}: {name} ({country}) - {reason}")

    print(f"\nOutput disimpan ke: {args.output}")
    print("Selesai!")


if __name__ == "__main__":
    main()
