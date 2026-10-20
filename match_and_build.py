import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
import re
import sys

OUTPUT_FILE = sys.argv[1]

# PDF files by country folder (from Google Drive listing)
pdf_files_by_country_folder = {
    "Singapore": [
        "Visa C2 2026  Muhammad Syawal bin Abdul Razak.pdf",
        "Visa C2 2025 Nicholas Gerard Gomez.pdf",
    ],
    "Maldives": [
        "Visa C2 2026 Mohamed Isham.pdf",
        "Visa C2 2026 Thamoor Adam Moosa.pdf",
        "Visa C2 2026 Aminath Maiha Ahmed.pdf",
        "Visa C2 2026 Chandana Dindu Kumara Hetti Arachchige.pdf",
    ],
    "Taiwan": [
        "Visa C2 2026 Chen Ning Chen.pdf",
        "Visa C2 2026 Kai Lun Huang.pdf",
        "Visa C2 2026 Shu Yuan Chen.pdf",
        "Visa C2 20226 Hao Pang Lin.pdf",
    ],
    "Malaysia": [
        "Visa C2 2026 Muhammad Faizal bin MD Fikri.pdf",
        "Visa C2 2026 Hafizul Affi bin Rusmani.pdf",
        "Visa C2 206 Aidil Adha bin Samsumar.pdf",
        "Visa C2 2026 Kwang Teck Hing.pdf",
        "Visa C2 2026 Mohd Nazri bin MD Yusoff.pdf",
        "Visa C2 2026 Muhammad Al Amin bin Yusri.pdf",
    ],
    "India": [
        "Visa C2 2026 Sagar Narayan Bhalerao.pdf",
        "Visa C2 2026 Vivek Saxena.pdf",
        "Visa C2 2026 Manoj Kumar.pdf",
        "Visa C2 2026 Maqbool Khan Abdulrazak.pdf",
        "Visa C2 2026 Amit Mahata.pdf",
        "Visa C2 2026 Gururaj  Vasantrao  Joshi.pdf",
        "Visa C2 2026 Aromal Parayankattuthara Robson.pdf",
        "Visa C2 2026 Ajay Dahiya.pdf",
    ],
    "Cambodia": [
        "Visa C2 2026 Sokon San.pdf",
        "Visa C2 2026 Dara Teng.pdf",
        "Visa C2 2026 Chetra Ry.pdf",
        "Visa C2 2026 Rattana Chhay.pdf",
        "Visa C2 2026 Rithy Seng.pdf",
        "Visa C2 2026 Sittichai Thaiking.pdf",
        "Visa C2 2026 Nakha Tuon.pdf",
        "Visa C2 2026 Makara Vorn.pdf",
    ],
    "Thailand": [
        "Visa C2 2026 Chawadol Reunthanakul.pdf",
        "Visa C2 2026 Tawatchai Praha.pdf",
        "Visa C2 2026 Pramote Inthisang.pdf",
        "Visa C2 2026 Nathawut Loisaioo.pdf",
        "Visa C2 2026 Wanchai Srinothai.pdf",
        "Visa C2 2026 Dusit Namsue.pdf",
        "Visa C2 2026 Tawee Bunrat.pdf",
        "Visa C2 2026 Thawin Jiengram.pdf",
        "Visa C2 2026 Tanasate Tanapornchockanun.pdf",
        "Visa C2 2026 Sitichok Reansritongkam.pdf",
        "Visa C2 2026 Noppharat Dimun.pdf",
    ],
    "Vietnam": [
        "Visa C2 2026 Quoc Long Mac.pdf",
        "Visa C2 2026 Xuan Bien Phung.pdf",
        "Visa C2 2026 Quang Hieu Dau.pdf",
        "Visa C2 2026 Duy Lam Vu.pdf",
        "Visa C2 2026 Van Phu Tran.pdf",
        "Visa C2 2026 Van Nguyen Nguyen.pdf",
        "Visa C2 2026 Thanh Tuan Le.pdf",
        "Visa C2 2026 Van Linh Nguyen.pdf",
        "Visa C2 2026 Cong Thanh Bui.pdf",
        "Visa C2 2026 Hong Quang Nguyen.pdf",
        "Visa C2 2025 Van Phu Tran.pdf",
        "Visa C2 2026 Xuan Thang Le.pdf",
        "Visa C2 2026 Van Tuan Song Hoang.pdf",
    ],
}

# Spreadsheet data (from read_file_content)
participants = [
    (1, "Mr.", "CHANDANA DINIDU KUMARA HETTI ARACHCHIGE", "Maldives", "Sheesha Pvt Ltd", "Contestant", "Batch 1"),
    (2, "Ms.", "AMINATH MAIHA AHMED", "Maldives", "Sheesha Pvt Ltd", "Escort", "Batch 1"),
    (3, "Mr.", "THAIMOOR ADAM MOOSA", "Maldives", "Sheesha Pvt Ltd", "Escort", "Batch 1"),
    (4, "Mr.", "MOHAMED ISHAM", "Maldives", "Sheesha Pvt Ltd", "Escort", "Batch 1"),
    (5, "Mr.", "RY CHETRA", "Cambodia", "NCX-C", "Contestant", "Batch 1"),
    (6, "Mr.", "SAN SOKON", "Cambodia", "NCX-C", "Contestant", "Batch 1"),
    (7, "Mr.", "TUON NAKHA", "Cambodia", "NCX-C", "Contestant", "Batch 1"),
    (8, "Mr.", "TENG DARA", "Cambodia", "NCX-C", "Contestant", "Batch 1"),
    (9, "Mr.", "VORN MAKARA", "Cambodia", "NCX-C", "Escort", "Batch 1"),
    (10, "Mr.", "SENG RITHY", "Cambodia", "NCX-C", "Escort", "Batch 1"),
    (11, "Mr.", "CHHAY RATTANA", "Cambodia", "NCX-C", "Escort", "Batch 1"),
    (12, "Mr.", "THAIKING SITTICHAI", "Cambodia", "NCX-C", "Escort", "Batch 1"),
    (13, "Mr.", "LIN HAO PANG", "Taiwan", "HTW", "Contestant", "Batch 1"),
    (14, "Mr.", "HUANG KAI LUN", "Taiwan", "HTW", "Contestant", "Batch 1"),
    (15, "Ms.", "CHEN CHIEH NING", "Taiwan", "HTW", "Genpo/ Dist", "Batch 1"),
    (16, "Mr.", "CHEN SU YUAN", "Taiwan", "HTW", "Genpo/ Dist", "Batch 1"),
    (17, "Mr.", "MUHAMMAD AL AMIN BIN YUSRI", "Malaysia", "BSH", "Contestant", "Batch 1"),
    (18, "Mr.", "KWANG TECK HING", "Malaysia", "BSH", "Contestant", "Batch 1"),
    (19, "Mr.", "AIDIL ADHA BIN SAMSUMAR", "Malaysia", "BSH", "Contestant", "Batch 1"),
    (20, "Mr.", "MUHAMMAD FAIZAL BIN MD FIKRI", "Malaysia", "BSH", "Escort", "Batch 1"),
    (21, "Mr.", "HAFIZUL AFIFI BIN RUSMANI", "Malaysia", "BSH", "Escort", "Batch 1"),
    (22, "Mr.", "MOHD NAZRI BIN MD YUSOFF", "Malaysia", "BSH", "Escort", "Batch 1"),
    (23, "Mr.", "NICHOLAS GERARD GOMEZ", "Singapore", "BSS", "Escort", "Batch 1"),
    (24, "Mr.", "MUHAMMAD SYAWAL BIN ABDUL RAZAK", "Singapore", "BSS", "Contestant", "Batch 1"),
    (25, "Mr.", "AMIT MAHATA", "India", "HMSI", "Contestant", "Batch 1"),
    (26, "Mr.", "MANOJ KUMAR", "India", "HMSI", "Contestant", "Batch 1"),
    (27, "Mr.", "AROMAL PARAYANKATTUTHARA ROBSON", "India", "HMSI", "Contestant", "Batch 1"),
    (28, "Mr.", "ABDULRAZAK MAQBOOL KHAN", "India", "HMSI", "Contestant", "Batch 1"),
    (29, "Mr.", "GURURAJ VASANTRAO JOSHI", "India", "HMSI", "Genpo/ Dist", "Batch 1"),
    (30, "Mr.", "AJAY DAHIYA", "India", "HMSI", "Genpo/ Dist", "Batch 1"),
    (31, "Mr.", "VIVEK SAXENA", "India", "HMSI", "Genpo/ Dist", "Batch 1"),
    (32, "Mr.", "SAGAR NARAYAN BHALERAO", "India", "HMSI", "Genpo/ Dist", "Batch 1"),
    (33, "Mr.", "PRASERT VICHITKULSAWAT", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (34, "Mr.", "PAKORN BOONPIAMSAK", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (35, "Ms.", "KANCHANA PARINYAPORNKUN", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (36, "Mr.", "HIROSHI ISHIDA", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (37, "Mr.", "TOSHIYA TAKENAGA", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (38, "Mr.", "WISIT SOMWANGSIRI", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (39, "Ms.", "KANLAYAPHAS BOONYATEERAPAT", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (40, "Ms.", "SUKRIS KONGPONG", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (41, "Ms.", "NUTCHAREE SUKHO", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (42, "Ms.", "MANIDA KAEWPHETSUWAN", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (43, "Ms.", "PHORNTHIP SUNTISUKKASEM", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (44, "Mr.", "TEERAPONG PROMDEE", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (45, "Ms.", "CHOMPOONUCH ASAWACHAISUWIKROM", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (46, "Mr.", "POOHRIN WUTTHIKAIUDOMDAJ", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (47, "Mr.", "CHANATHIP SARTTAYARAT", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (48, "Ms.", "NAPHATSORN SILWATTA", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (49, "Ms.", "WIMONRAT HIRANWONG", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (50, "Mr.", "PONGSATORN PROMWANG", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (51, "Ms.", "PICHAYAPORN CHAIYARUNGROJ", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (52, "Ms.", "SARISA INGKANANTAWAREE", "Thailand", "ASH", "ASH Committee", "Batch 1"),
    (53, "Mr.", "Hoang Van Tuan Song", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (54, "Mr.", "Le Xuan Thang", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (55, "Mr.", "Nguyen Van Linh", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (56, "Mr.", "Tran Van Phu", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (57, "Mr.", "Le Thanh Tuan", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (58, "Mr.", "Mac Quoc Long", "Vietnam", "HVN", "Contestant", "Batch 1"),
    (59, "Mr.", "Tran Bao Kha", "Vietnam", "HVN", "Dealer", "Batch 1"),
    (60, "Mr.", "Bui Cong Thanh", "Vietnam", "HVN", "Dealer", "Batch 1"),
    (61, "Mr.", "Nguyen Hong Quang", "Vietnam", "HVN", "Dealer", "Batch 1"),
    (62, "Mr.", "Nguyen Hong Thanh", "Vietnam", "HVN", "Dealer", "Batch 1"),
    (63, "Mr.", "Vu Duy Lam", "Vietnam", "HVN", "Escort", "Batch 1"),
    (64, "Mr.", "Dau Quang Hieu", "Vietnam", "HVN", "Escort", "Batch 1"),
    (65, "Mr.", "Nguyen Van Chuyen", "Vietnam", "HVN", "Escort", "Batch 1"),
    (66, "Mr.", "Phung Xuan Bien", "Vietnam", "HVN", "Genpo/ Dist", "Batch 1"),
    (67, "Mr.", "AARON GLENN P. AGUILAR", "Philippines", "HPI", "Contestant", "Batch 2"),
    (68, "Mr.", "JAY M. BERNABE", "Philippines", "HPI", "Contestant", "Batch 2"),
    (69, "Mr.", "ALBERTO M. ALABAR", "Philippines", "HPI", "Contestant", "Batch 2"),
    (70, "Mr.", "JUSTINE B. SUYOM", "Philippines", "HPI", "Contestant", "Batch 2"),
    (71, "Mr.", "MARK JOHN V. BLAS", "Philippines", "HPI", "Contestant", "Batch 2"),
    (72, "Mr.", "PETER JHON M. BARENG", "Philippines", "HPI", "Genpo/ Dist", "Batch 2"),
    (73, "Mr.", "CYRUS L. PAGARAN", "Philippines", "HPI", "Genpo/ Dist", "Batch 2"),
    (74, "Mr.", "KEIZO OTA", "Thailand", "ASH", "ASH Committee", "Batch 2"),
    (75, "Mr.", "TAWATCHAI PRAHA", "Thailand", "TH", "Genpo/ Dist", "Batch 2"),
    (76, "Mr.", "DUSIT NAMSUE", "Thailand", "TH", "Escort", "Batch 2"),
    (77, "Mr.", "CHAWADOL REUNTHANAKUL", "Thailand", "TH", "Escort", "Batch 2"),
    (78, "Mr.", "TANASATE TANAPORNCHOCKANUN", "Thailand", "TH", "Escort", "Batch 2"),
    (79, "Mr.", "NOPPHARAT DIMUN", "Thailand", "TH", "Contestant", "Batch 2"),
    (80, "Mr.", "TAWEE BUNRAT", "Thailand", "TH", "Contestant", "Batch 2"),
    (81, "Mr.", "WANCHAI SRINOTHAI", "Thailand", "TH", "Contestant", "Batch 2"),
    (82, "Mr.", "SITTICHOK REANSRITONGKAM", "Thailand", "TH", "Contestant", "Batch 2"),
    (83, "Mr.", "THAWIN JIENGRAM", "Thailand", "TH", "Contestant", "Batch 2"),
    (84, "Mr.", "NATHAWUT LOISAIOO", "Thailand", "TH", "Contestant", "Batch 2"),
    (85, "Mr.", "PRAMOTE INTHISANG", "Thailand", "TH", "Dealer", "Batch 2"),
    (86, "Ms", "LE THI THUY TRANG", "Vietnam", "HVN", "Dealer", "Batch 2"),
]

def extract_name_from_pdf(filename):
    name = filename.replace(".pdf", "")
    name = re.sub(r'^Visa C2\s+\d+\s+', '', name)
    return name.strip()

def normalize_name(name):
    name = name.lower().strip()
    name = re.sub(r'\s+', ' ', name)
    return name

def name_words(name):
    n = normalize_name(name)
    return set(n.split())

def levenshtein(s1, s2):
    if len(s1) < len(s2):
        return levenshtein(s2, s1)
    if len(s2) == 0:
        return len(s1)
    prev_row = range(len(s2) + 1)
    for i, c1 in enumerate(s1):
        curr_row = [i + 1]
        for j, c2 in enumerate(s2):
            insertions = prev_row[j + 1] + 1
            deletions = curr_row[j] + 1
            substitutions = prev_row[j] + (c1 != c2)
            curr_row.append(min(insertions, deletions, substitutions))
        prev_row = curr_row
    return prev_row[-1]

def word_similarity(w1, w2):
    if w1 == w2:
        return 1.0
    if w1 in w2 or w2 in w1:
        return max(len(w1), len(w2)) / (max(len(w1), len(w2)) + 1)
    d = levenshtein(w1, w2)
    max_len = max(len(w1), len(w2))
    return 1.0 - d / max_len if max_len > 0 else 0.0

def fuzzy_word_match(name1_words, name2_words):
    if not name1_words or not name2_words:
        return 0.0
    common = name1_words & name2_words
    total = name1_words | name2_words
    if len(common) == len(total):
        return 1.0
    unmatched1 = name1_words - common
    unmatched2 = name2_words - common
    fuzzy_matches = 0
    used2 = set()
    for w1 in unmatched1:
        best_score = 0
        best_w2 = None
        for w2 in unmatched2:
            if w2 in used2:
                continue
            score = word_similarity(w1, w2)
            if score > best_score:
                best_score = score
                best_w2 = w2
        if best_score >= 0.65 and best_w2:
            fuzzy_matches += 1
            used2.add(best_w2)
    exact_weight = len(common) * 1.0
    fuzzy_weight = fuzzy_matches * 0.7
    total_words = max(len(name1_words), len(name2_words))
    return (exact_weight + fuzzy_weight) / total_words if total_words > 0 else 0.0

# Build PDF name index by country
pdf_names_by_country = {}
for folder_country, files in pdf_files_by_country_folder.items():
    names = []
    for f in files:
        name = extract_name_from_pdf(f)
        names.append((name, name_words(name)))
    pdf_names_by_country[folder_country.lower()] = names

# Match each participant
results = []
matched_count = 0
unmatched_count = 0

for no, title, name, country, company, role, batch in participants:
    country_key = country.lower()
    pdf_names = pdf_names_by_country.get(country_key, [])

    if not pdf_names:
        results.append({
            "no": no, "title": title, "name": name, "country": country,
            "company": company, "role": role, "batch": batch,
            "match": None, "score": 0,
            "status": "NO_FOLDER" if country_key == "philippines" else "NO_PDFS",
            "check": "",
        })
        unmatched_count += 1
        continue

    spreadsheet_words = name_words(name)
    best_score = 0
    best_match = None

    for pdf_name, pdf_words in pdf_names:
        score = fuzzy_word_match(spreadsheet_words, pdf_words)
        if score > best_score:
            best_score = score
            best_match = pdf_name

    if best_score >= 0.7:
        matched_count += 1
        results.append({
            "no": no, "title": title, "name": name, "country": country,
            "company": company, "role": role, "batch": batch,
            "match": best_match, "score": round(best_score, 2),
            "status": "MATCHED", "check": "√",
        })
    else:
        unmatched_count += 1
        results.append({
            "no": no, "title": title, "name": name, "country": country,
            "company": company, "role": role, "batch": batch,
            "match": best_match, "score": round(best_score, 2),
            "status": "NOT_FOUND", "check": "",
        })

# Build xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"

headers = ["No.", "Title", "Full Name", "Country", "Company", "Participant Role", "Batch", "Ket"]
header_font = Font(bold=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')
    cell.border = thin_border

for row_idx, r in enumerate(results, 2):
    ws.cell(row=row_idx, column=1, value=r["no"]).border = thin_border
    ws.cell(row=row_idx, column=2, value=r["title"]).border = thin_border
    ws.cell(row=row_idx, column=3, value=r["name"]).border = thin_border
    ws.cell(row=row_idx, column=4, value=r["country"]).border = thin_border
    ws.cell(row=row_idx, column=5, value=r["company"]).border = thin_border
    ws.cell(row=row_idx, column=6, value=r["role"]).border = thin_border
    ws.cell(row=row_idx, column=7, value=r["batch"]).border = thin_border
    ket_cell = ws.cell(row=row_idx, column=8, value=r["check"])
    ket_cell.alignment = Alignment(horizontal='center')
    ket_cell.border = thin_border

# Column widths
ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 6
ws.column_dimensions['C'].width = 45
ws.column_dimensions['D'].width = 15
ws.column_dimensions['E'].width = 20
ws.column_dimensions['F'].width = 18
ws.column_dimensions['G'].width = 10
ws.column_dimensions['H'].width = 8

# Add legend row
legend_row = len(results) + 3
ws.cell(row=legend_row, column=1, value="Ketarangan :")
ws.cell(row=legend_row + 1, column=1, value="√ sudah selesai")

wb.save(OUTPUT_FILE)

# Print results
print(f"\n{'='*100}")
print(f"VISA C2 MATCHING RESULTS")
print(f"{'='*100}")
print(f"Total participants: {len(results)}")
print(f"Matched (√): {matched_count}")
print(f"Not matched: {unmatched_count}")
print(f"{'='*100}\n")

print(f"{'No.':<5} {'Status':<12} {'Score':<7} {'Spreadsheet Name':<45} {'Country':<15} {'Matched PDF Name'}")
print(f"{'-'*5} {'-'*12} {'-'*7} {'-'*45} {'-'*15} {'-'*40}")
for r in results:
    status_icon = "V" if r['status'] == 'MATCHED' else "X"
    match_str = r['match'] or "---"
    print(f"{r['no']:<5} {status_icon} {r['status']:<10} {r['score']:<7} {r['name']:<45} {r['country']:<15} {match_str}")

print(f"\n{'='*100}")
print(f"UNMATCHED PARTICIPANTS:")
print(f"{'='*100}")
for r in results:
    if r['status'] != 'MATCHED':
        print(f"  No.{r['no']} - {r['name']} ({r['country']}) - Status: {r['status']}, Best match: {r['match']} (score: {r['score']})")
